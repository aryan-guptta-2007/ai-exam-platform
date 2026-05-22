from typing import List, Dict, Any
from app.utils.parsers.base import BaseParser
from loguru import logger

class PPTXParser(BaseParser):
    def parse(self, file_path: str) -> List[Dict[str, Any]]:
        logger.info(f"Parsing PPTX file: {file_path}")
        results = []
        
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            
            for i, slide in enumerate(prs.slides):
                slide_num = i + 1
                slide_texts = []
                
                for shape in slide.shapes:
                    if hasattr(shape, "text_frame") and shape.text_frame:
                        for paragraph in shape.text_frame.paragraphs:
                            if paragraph.text.strip():
                                slide_texts.append(paragraph.text.strip())
                                
                slide_content = "\n".join(slide_texts)
                cleaned_content = slide_content.strip()
                
                results.append({
                    "content": cleaned_content,
                    "page_number": slide_num,
                    "char_start": 0,
                    "char_end": len(cleaned_content),
                    "paragraph_references": slide_texts
                })
                
            logger.info(f"Successfully extracted {len(results)} slides from PPTX.")
            return results
            
        except Exception as e:
            logger.error(f"Error parsing PPTX file {file_path}: {e}")
            raise e
